"""Construye el reporte ejecutivo .docx a partir de resultados.json, tablas/ y figuras/."""
import re
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor
import respuestas as RS

R = RS.R; fm = RS.fm; n1, n2, n3 = RS.n1, RS.n2, RS.n3
AZUL = RGBColor(0x1F, 0x4E, 0x79)
doc = Document()

# ---------- página y estilos ----------
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for m in ("left_margin", "right_margin"): setattr(sec, m, Inches(1))
sec.top_margin = sec.bottom_margin = Inches(0.9)
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.1
for nombre, tam, color in [("Heading 1", 16, AZUL), ("Heading 2", 13, AZUL), ("Heading 3", 11, AZUL)]:
    h = doc.styles[nombre]; h.font.name = "Calibri"; h.font.size = Pt(tam); h.font.bold = True; h.font.color.rgb = color
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri"); h.element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); h.element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    h.paragraph_format.space_before = Pt(14 if nombre == "Heading 1" else 10); h.paragraph_format.space_after = Pt(6); h.paragraph_format.keep_with_next = True

def runs(p, texto, size=None, color=None):
    """Agrega texto con **negritas** y *cursivas* simples."""
    texto = texto.replace("`", "")
    for parte in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*)", texto):
        if not parte: continue
        if parte.startswith("**"): r = p.add_run(parte[2:-2]); r.bold = True
        elif parte.startswith("*"): r = p.add_run(parte[1:-1]); r.italic = True
        else: r = p.add_run(parte)
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
    return p
def para(texto, align=None, size=None, color=None, after=None, indent=None):
    p = doc.add_paragraph(); runs(p, texto, size, color)
    if align: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if indent: p.paragraph_format.left_indent = Inches(indent)
    return p
def sombra(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hex_); tcPr.append(sh)
def tabla(df, anchos=None, size=8.5, titulo=None, fuente_nota=None):
    if titulo:
        p = para(f"**{titulo}**", size=9.5, after=3); p.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(df.columns)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, c in enumerate(df.columns):
        cel = t.rows[0].cells[j]; cel.text = ""; p = cel.paragraphs[0]; r = p.add_run(str(c)); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER; sombra(cel, "1F4E79")
    for i, fila in enumerate(df.itertuples(index=False)):
        cs = t.add_row().cells
        for j, v in enumerate(fila):
            cs[j].text = ""; p = cs[j].paragraphs[0]; r = p.add_run(str(v)); r.font.size = Pt(size)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 or isinstance(v, str) and not re.match(r"^([-+−]?[\d.,]+%?|NaN)$", v) else WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.space_after = Pt(0)
            if i % 2: sombra(cs[j], "EAF1F8")
    if anchos:
        for row in t.rows:
            for j, w in enumerate(anchos): row.cells[j].width = Inches(w)
    # repetir encabezado y no partir filas
    for k_, row in enumerate(t.rows):                      # orden del esquema: cantSplit antes que tblHeader
        trPr = row._tr.get_or_add_trPr(); cs_ = OxmlElement("w:cantSplit"); cs_.set(qn("w:val"), "true"); trPr.append(cs_)
        if k_ == 0: th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    if len(t.rows) <= 14:
        for row in t.rows[:-1]:
            for cel in row.cells:
                for pp in cel.paragraphs: pp.paragraph_format.keep_with_next = True
    if fuente_nota: para(f"*{fuente_nota}*", size=8.5, after=8)
    else: doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t
cont_fig = [0]
def figura(archivo, ancho, pie, interp):
    cont_fig[0] += 1
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(f"figuras/{archivo}.png", width=Inches(ancho))
    c = para(f"*Figura {cont_fig[0]}. {pie}*", align=WD_ALIGN_PARAGRAPH.CENTER, size=9, after=3)
    para(f"**Interpretación.** {interp}", size=10)
def pregs(s, pie_titulo="Respuestas a las preguntas de la sección"):
    doc.add_heading(pie_titulo, level=3)
    for q, a in RS.Q[s]:
        p = para(f"**{q}**", after=1); p.paragraph_format.keep_with_next = True
        para(a, indent=0.2)
def csv(nombre): return pd.read_csv(f"tablas/{nombre}.csv")
def f(x, d=2): return f"{x:,.{d}f}"
def pagina_nueva(): doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ---------- pie de página con numeración ----------
def campo(par, instr):
    for tipo in ("begin", "instr", "separate", "texto", "end"):
        r = par.add_run(); r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        if tipo == "instr":
            i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = f" {instr} "; r._r.append(i)
        elif tipo == "texto": r.text = "1"
        else:
            c = OxmlElement("w:fldChar"); c.set(qn("w:fldCharType"), tipo); r._r.append(c)
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
rr = fp.add_run("Series de Tiempo Aplicadas a las Finanzas · Trabajo Final Unidad 1 · Rocket Lab (RKLB) · Página "); rr.font.size = Pt(8.5); rr.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
campo(fp, "PAGE")
for r_ in fp.runs: r_.font.size = Pt(8.5)

# ======================= PORTADA =======================
for _ in range(4): doc.add_paragraph()
para("TRABAJO FINAL · UNIDAD 1", align=WD_ALIGN_PARAGRAPH.CENTER, size=13, color=AZUL, after=4)
p = para("**Diagnóstico y preparación de una serie de tiempo financiera**", align=WD_ALIGN_PARAGRAPH.CENTER, size=26, color=AZUL, after=6)
para("Precio de cierre mensual de Rocket Lab USA, Inc. (Nasdaq: RKLB)", align=WD_ALIGN_PARAGRAPH.CENTER, size=15, after=40)
portada = pd.DataFrame({"": ["Materia", "Unidad", "Serie analizada", "Periodo / frecuencia", "Integrantes", "Docente", "Fecha de entrega"],
                        " ": ["Series de Tiempo Aplicadas a las Finanzas", "Unidad 1 · Diagnóstico y preparación de una serie de tiempo",
                              "Rocket Lab USA, Inc. (RKLB) · precio de cierre, USD por acción", f"{fm(R['ini'])} a {fm(R['fin'])} · mensual · {R['n']} observaciones",
                              "[Completar nombres de los integrantes]", "[Completar]", "[Completar fecha]"]})
t = doc.add_table(rows=0, cols=2); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
for a, b in portada.itertuples(index=False):
    cs = t.add_row().cells; cs[0].text = ""; cs[1].text = ""
    r0 = cs[0].paragraphs[0].add_run(a); r0.bold = True; r0.font.size = Pt(11); r0.font.color.rgb = RGBColor(255, 255, 255); sombra(cs[0], "1F4E79")
    r1 = cs[1].paragraphs[0].add_run(b); r1.font.size = Pt(11)
    cs[0].width = Inches(1.7); cs[1].width = Inches(4.6)
doc.add_paragraph()
para("Documento que acompaña al notebook `RKLB_Trabajo_Final_Unidad1.ipynb`. Fecha de consulta de los datos: 2 de octubre de 2026.", align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5, color=RGBColor(0x59, 0x59, 0x59))
pagina_nueva()

# ======================= SÍNTESIS =======================
doc.add_heading("Síntesis ejecutiva", level=1)
para(f"Se analizó el precio de cierre mensual de Rocket Lab (RKLB) entre {fm(R['ini'])} y {fm(R['fin'])} ({R['n']} observaciones, USD por acción). La base está completa: no tiene faltantes, duplicados ni meses ausentes; el único detalle es que el último mes ({fm(R['fin'])}) está incompleto a la fecha de consulta.")
for b in [
    f"**Tendencia y escala.** El precio pasó de {n2(R['primero'])} a {n2(R['ultimo'])} USD (+{R['simple_acum']:,.0f} %), con un máximo de {n2(R['maximo'])} USD en {fm(R['f_max'])} y una corrección posterior de {abs(R['caida_max']):.0f} %. El crecimiento es casi exponencial (≈ {n1(R['g_mensual'])} % mensual compuesto) y la variabilidad crece con el nivel.",
    f"**Transformaciones.** El logaritmo solo reescala (autocorrelación de rezago 1 = 0.977); la primera diferencia quita la tendencia pero conserva la heterocedasticidad; el **cambio logarítmico** deja una serie sin tendencia visible, con autocorrelación de rezago 1 = {n3(R['lag1_ret'])} y media de {n2(R['r_media'])} % mensual (desv. {n1(R['r_std'])} %).",
    f"**Dependencia temporal.** En nivel, la ACF es {n3(RS.A('acf_nivel',1,'ACF'))} en el rezago 1, con los rezagos 1–4 fuera de bandas y decaimiento lento; la PACF concentra la dependencia en el rezago 1. En el cambio logarítmico ninguna barra de ACF ni de PACF supera las bandas (Ljung-Box p = {n3(R['lb_r_p'])}).",
    f"**Episodios.** nov-2024 (+154.95 % en un mes) es el movimiento más extremo; hay rachas de alzas (may–nov-2024, abr–ago-2025), un posible cambio de nivel tras nov-2024 y regímenes de volatilidad (desv. móvil de 6 meses entre {n1(R['vm_min'])} % y {n1(R['vm_max'])} %).",
    f"**Train/Test.** División cronológica 80/20: Train de {fm(R['tr_ini'])} a {fm(R['tr_fin'])} ({R['n_train']} meses) y Test de {fm(R['te_ini'])} a {fm(R['te_fin'])} ({R['n_test']} meses).",
    f"**Representación recomendada para modelar:** el cambio logarítmico mensual, estimado con Train; se sugiere considerar un modelo de volatilidad condicional. Limitación principal: muestra corta ({R['n']-1} cambios), que hace amplias las bandas de confianza.",
]:
    p = doc.add_paragraph(style="List Bullet"); runs(p, b); p.paragraph_format.space_after = Pt(3)
rt = {r["Transformación"]: r for r in R["res_transf"]}
resumen = pd.DataFrame([
    ["Observaciones", f"{R['n']} meses ({fm(R['ini'])} – {fm(R['fin'])})"],
    ["Precio inicial → final (USD)", f"{n2(R['primero'])} → {n2(R['ultimo'])}"],
    ["Máximo / mínimo (USD)", f"{n2(R['maximo'])} ({fm(R['f_max'])}) / {n2(R['minimo'])} ({fm(R['f_min'])})"],
    ["Autocorr. rezago 1: nivel / ln / ΔY / cambio log", f"{rt['Nivel (USD)']['autocorr. lag 1']:.3f} / {rt['ln(Y)']['autocorr. lag 1']:.3f} / {rt['1ª diferencia (USD)']['autocorr. lag 1']:.3f} / {rt['Cambio log (%)']['autocorr. lag 1']:.3f}"],
    ["Rezagos fuera de bandas: ACF nivel / PACF nivel", f"{RS.R['acf_nivel_fuera']} / {RS.R['pacf_nivel_fuera']}"],
    ["Rezagos fuera de bandas: ACF / PACF del cambio log", "ninguno / ninguno"],
    ["Cambio log mensual: media / desv. est.", f"{n2(R['r_media'])} % / {n2(R['r_std'])} %"],
    ["Train / Test", f"{R['n_train']} / {R['n_test']} observaciones"]], columns=["Indicador", "Valor"])
tabla(resumen, anchos=[2.9, 3.6], size=9.5, titulo="Tabla 1. Resumen de resultados clave")

# ======================= 1 =======================
doc.add_heading("1. Descripción de la serie", level=1)
para("La serie se tomó del archivo de datos históricos de RKLB proporcionado por el equipo. La tabla resume su identificación; el archivo original y la versión limpia (`RKLB_base_limpia.csv`) se entregan junto con este reporte.")
tabla(csv("t01_identificacion"), anchos=[1.8, 4.7], size=9.5, titulo="Tabla 2. Identificación de la serie")
pregs(1)

# ======================= 2 =======================
doc.add_heading("2. Preparación de datos", level=1)
chk = pd.DataFrame([
    ["Valores faltantes (7 columnas)", f"{R['nulos']}", "Sin acción"],
    ["Filas duplicadas / meses repetidos", f"{R['dup_filas']} / {R['dup_fechas']}", "Sin acción"],
    ["Meses del calendario vs. meses con dato", f"{R['n_meses_cal']} vs. {R['n']}", "Serie continua, sin huecos"],
    ["Etiquetas de mes en fin de semana", f"{R['n_finsem']} de {R['n']}", "Son etiquetas de mes (día 1); no se modifican"],
    ["Orden de la fecha en el archivo", "Descendente", "Se ordenó ascendente"],
    ["«% var.» fuente vs. cálculo con el cierre", f"≤ {R['dif_max_pct']:.4f} pp", "Consistente (redondeo)"],
    ["Mes parcial (oct-2026)", f"Vol. {n2(R['vol_ult'])} M = {R['vol_ult']/R['vol_med_resto']*100:.1f} % de la mediana", "Se conserva; se verifica sensibilidad"]],
    columns=["Revisión", "Resultado", "Decisión"])
tabla(chk, anchos=[2.6, 1.9, 2.0], size=9, titulo="Tabla 3. Revisión de calidad de los datos")
pregs(2)

# ======================= 3 =======================
doc.add_heading("3. Serie en nivel", level=1)
d3 = csv("t03_descriptiva_nivel").rename(columns={"Unnamed: 0": ""}).drop(columns=[""]).iloc[0]
t4v = pd.DataFrame({"Estadístico": ["Observaciones (n)", "Media", "Desviación estándar", "Mínimo", "Primer cuartil (Q1)", "Mediana", "Tercer cuartil (Q3)", "Máximo", "Rango", "Varianza", "Coeficiente de variación (%)", "Asimetría", "Curtosis (exceso)"],
                    "Valor": [f"{int(d3['n'])}"] + [f"{d3[k]:,.2f}" for k in ["media", "desv. est.", "mínimo", "Q1", "mediana", "Q3", "máximo", "rango", "varianza", "coef. variación (%)", "asimetría", "curtosis (exceso)"]]})
tabla(t4v, anchos=[2.6, 1.5], size=9, titulo="Tabla 4. Estadísticos descriptivos del precio de cierre (USD por acción)", fuente_nota="Fuente: cálculos propios con la base RKLB (n = 49 meses). La media (29.87) casi triplica la mediana (10.70) y el coeficiente de variación es de 109 %: distribución muy asimétrica (1.35) por el auge reciente.")
figura("fig_03_nivel", 6.5, "Precio de cierre mensual de RKLB en nivel (izquierda: escala lineal; derecha: escala logarítmica).",
       f"El precio permanece entre 3.76 y 7.37 USD hasta mediados de 2024 y luego sube con saltos hasta {n2(R['maximo'])} USD en {fm(R['f_max'])}; después cae a ≈ 65–70 USD. La recta de tendencia (+{n2(R['pend'])} USD/mes, R² = {n3(R['r2'])}) ajusta peor que una tendencia exponencial (R² = {n3(R['r2_log'])} en logaritmos). En escala logarítmica las oscilaciones se ven más parejas, lo que confirma que la variabilidad crece con el nivel.")
po = pd.DataFrame(R["por_anio"]).rename(columns={"anio": "Año", "n": "n", "media": "Media (USD)", "desv_est": "Desv. est. (USD)", "minimo": "Mín. (USD)", "maximo": "Máx. (USD)", "coef_var_%": "Coef. var. (%)"})
po["Año"] = po["Año"].astype(int).astype(str); po["n"] = po["n"].astype(int)
tabla(po.astype(object).apply(lambda c: c if c.name == "Año" else (c.map(lambda v: f"{int(v)}") if c.name == "n" else c.map(lambda v: f"{v:,.2f}"))), size=9, titulo="Tabla 5. Precio de cierre por año calendario", fuente_nota="2022 incluye solo oct–dic; 2026 incluye ene–oct (oct parcial). Interpretación: la desviación estándar pasa de 0.67 a 24.82 USD, evidencia de variabilidad no constante.")
pregs(3)

# ======================= 4 =======================
doc.add_heading("4. Primera diferencia (ΔYₜ = Yₜ − Yₜ₋₁)", level=1)
t4 = csv("t04_primera_diferencia").iloc[:8].copy(); t4.columns = ["Mes", "Cierre (USD)", "ΔY (USD)"]
t4["Cierre (USD)"] = t4["Cierre (USD)"].map(lambda v: f"{v:,.2f}"); t4["ΔY (USD)"] = t4["ΔY (USD)"].map(lambda v: "NaN" if pd.isna(v) else f"{v:+,.2f}")
tabla(t4, anchos=[1.3, 1.5, 1.5], size=9, titulo="Tabla 6. Primera diferencia (primeras 8 observaciones; la tabla completa está en el notebook)")
figura("fig_04_primera_diferencia", 6.3, "Primera diferencia del precio de cierre mensual de RKLB (USD por acción).",
       f"La serie oscila alrededor de {RS.sg2(R['d_media'])} USD (sin tendencia), pero las barras son diminutas hasta 2024 y gigantes desde 2025: el mayor aumento es {RS.sg2(R['d_max'])} USD en {fm(R['f_d_max'])} y la mayor caída {RS.sg2(R['d_min'])} USD en {fm(R['f_d_min'])}. La variabilidad creciente sigue presente.")
pregs(4)

# ======================= 5 =======================
doc.add_heading("5. Logaritmo natural", level=1)
t5 = csv("t05_log_nivel").iloc[:6].copy(); t5.columns = ["Mes", "Cierre (USD)", "ln(Cierre)"]
t5["Cierre (USD)"] = t5["Cierre (USD)"].map(lambda v: f"{v:,.2f}"); t5["ln(Cierre)"] = t5["ln(Cierre)"].map(lambda v: f"{v:.4f}")
tabla(t5, anchos=[1.3, 1.5, 1.5], size=9, titulo="Tabla 7. Logaritmo natural del cierre (primeras 6 observaciones)")
figura("fig_05_log_nivel", 6.5, "Nivel (izquierda) y logaritmo natural (derecha) del precio de cierre mensual de RKLB.",
       f"El logaritmo conserva la trayectoria creciente (no la elimina) pero comprime el rango: de un máximo {R['razon_max_min']:.1f} veces el mínimo a una amplitud de {n2(R['ln_max']-R['ln_min'])} unidades. Las oscilaciones se ven más parejas que en nivel, aunque la serie sigue subiendo.")
pregs(5)

# ======================= 6 =======================
doc.add_heading("6. Cambio porcentual simple", level=1)
t6 = csv("t06_cambio_simple").iloc[:8].copy(); t6.columns = ["Mes", "Cierre (USD)", "ΔY (USD)", "Cambio simple (%)"]
t6["Cierre (USD)"] = t6["Cierre (USD)"].map(lambda v: f"{v:,.2f}")
t6["ΔY (USD)"] = t6["ΔY (USD)"].map(lambda v: "NaN" if pd.isna(v) else f"{v:+,.2f}"); t6["Cambio simple (%)"] = t6["Cambio simple (%)"].map(lambda v: "NaN" if pd.isna(v) else f"{v:+,.2f}")
tabla(t6, anchos=[1.1, 1.3, 1.3, 1.5], size=9, titulo="Tabla 8. Cambio porcentual simple (primeras 8 observaciones; la primera es NaN)")
figura("fig_06_cambio_simple", 6.3, "Cambio porcentual simple mensual del cierre de RKLB (%).",
       f"La serie oscila alrededor de cero sin tendencia y con oscilaciones de amplitud similar entre 2023 y 2026 (a diferencia de ΔY). Hay {R['p_pos']} meses positivos y {R['p_neg']} negativos; el extremo es {RS.sg2(R['p_max'])} % en {fm(R['f_p_max'])} y la mayor caída {RS.sg2(R['p_min'])} % en {fm(R['f_p_min'])}. La asimetría (alzas hasta +155 %, caídas hasta −36 %) es propia del cambio simple.")
pregs(6)

# ======================= 7 =======================
doc.add_heading("7. Cambio logarítmico", level=1)
t7 = csv("t07_cambio_log").iloc[:8].copy(); t7.columns = ["Mes", "Cierre (USD)", "Cambio simple (%)", "Cambio log (%)"]
t7["Cierre (USD)"] = t7["Cierre (USD)"].map(lambda v: f"{v:,.2f}")
for c in ["Cambio simple (%)", "Cambio log (%)"]: t7[c] = t7[c].map(lambda v: "NaN" if pd.isna(v) else f"{v:+,.2f}")
tabla(t7, anchos=[1.1, 1.3, 1.5, 1.5], size=9, titulo="Tabla 9. Cambio logarítmico vs. cambio simple (primeras 8 observaciones)")
figura("fig_07_cambio_log", 6.2, "Arriba: cambio logarítmico mensual del cierre de RKLB. Abajo: comparación con el cambio simple (%).",
       f"El cambio logarítmico oscila alrededor de {n2(R['r_media'])} % sin tendencia. En el panel inferior, las dos curvas casi coinciden en los meses tranquilos y se separan en los extremos (nov-2024: +154.95 % simple vs +93.59 % logarítmico), porque el logaritmo comprime las variaciones grandes.")
pregs(7)

# ======================= 8 =======================
doc.add_heading("8. Comparación de transformaciones", level=1)
t8 = csv("t08_comparacion_completa").iloc[:10].copy()
fmt8 = {"Nivel (USD)": "{:,.2f}", "1ª diferencia (USD)": "{:+,.2f}", "ln(Y)": "{:.4f}", "Cambio simple (%)": "{:+,.2f}", "Cambio log (%)": "{:+,.2f}"}
t8.columns = ["Mes"] + list(t8.columns[1:])
for c, fmt_ in fmt8.items(): t8[c] = t8[c].map(lambda v, fmt_=fmt_: "NaN" if pd.isna(v) else fmt_.format(v))
tabla(t8, size=8.5, titulo="Tabla 10. Las cinco representaciones de la serie (primeros 10 meses)", fuente_nota="La tabla completa de 49 meses está en el notebook y en tablas/t08_comparacion_completa.csv.")
s8 = pd.DataFrame(R["res_transf"]).drop(columns=["n"]) if False else pd.DataFrame(R["res_transf"])
s8["n"] = s8["n"].astype(int).astype(str)
for c in ["media", "desv. est.", "mín", "máx"]: s8[c] = s8[c].map(lambda v: f"{v:,.2f}")
for c in ["autocorr. lag 1", "autocorr. lag 6"]: s8[c] = s8[c].map(lambda v: f"{v:.3f}")
s8["p-valor ADF"] = s8["p-valor ADF"].map(lambda v: "< 0.001" if v < 0.001 else f"{v:.3f}")
tabla(s8, size=8.5, titulo="Tabla 11. Resumen comparativo de las transformaciones",
      fuente_nota="El p-valor de ADF es solo orientativo (muestra pequeña): p alto = no se rechaza raíz unitaria. Nota: la autocorr. de rezago 6 de ΔY (−0.473) está dominada por los saltos de 2026.")
figura("fig_08_comparacion", 6.1, "Panel con las cinco representaciones de la serie de RKLB.",
       "El nivel y ln(Y) comparten tendencia ascendente; ΔY elimina la tendencia pero con oscilaciones crecientes; los dos cambios relativos (simple y logarítmico) oscilan alrededor de cero con amplitud comparable en todo el periodo, por lo que son las representaciones que más reducen la tendencia y la persistencia visual.")
pregs(8)

# ======================= 9 =======================
doc.add_heading("9. División Train/Test", level=1)
sp = csv("t09_train_test"); sp["% de la muestra"] = sp["% de la muestra"].map(lambda v: f"{v:.1f}"); sp["Precio medio (USD)"] = sp["Precio medio (USD)"].map(lambda v: f"{v:,.2f}")
sp["Desv. est. cambio log (%)"] = sp["Desv. est. cambio log (%)"].map(lambda v: f"{v:,.2f}")
tabla(sp, size=8.5, titulo="Tabla 12. División cronológica 80 % / 20 %")
figura("fig_09_train_test", 6.3, "División cronológica Train/Test del precio de cierre mensual de RKLB (línea punteada: punto de corte).",
       f"Train ({R['n_train']} meses) cubre casi todo el periodo de precios bajos y la primera parte del auge; Test ({R['n_test']} meses) contiene el máximo de {n2(R['max_test'])} USD y la posterior corrección. El nivel de Test queda en su mayor parte fuera del rango de Train (media {n2(R['media_test'])} vs {n2(R['media_train'])} USD), lo que motiva modelar cambios relativos en lugar de niveles.")
pregs(9)

# ======================= 10 =======================
doc.add_heading("10. Rezagos y autocorrelación", level=1)
lg = csv("t10_rezagos").head(8)
for c in lg.columns[1:]: lg[c] = lg[c].map(lambda v: "NaN" if pd.isna(v) else f"{v:,.2f}")
lg.columns = ["Mes", "Yₜ (USD)", "Lag 1", "Lag 2", "Lag 3"]
tabla(lg, size=9, titulo="Tabla 13. Rezagos del precio de cierre (primeras 8 filas; NaN donde no existe el pasado)")
a1 = csv("t10_autocorrelaciones_nivel"); a1.columns = ["Rezago (meses)", "Pares usados", "Correlación Pearson Yₜ vs Yₜ₋ₖ", "ACF muestral"]
a1[a1.columns[2]] = a1[a1.columns[2]].map(lambda v: f"{v:.4f}"); a1["ACF muestral"] = a1["ACF muestral"].map(lambda v: f"{v:.4f}")
tabla(a1, size=9, titulo="Tabla 14. Autocorrelaciones de la serie en nivel (rezagos 1 a 3)")
a2 = csv("t10_autocorrelaciones_retorno"); a2.columns = ["Rezago (meses)", "Correlación Pearson", "ACF muestral"]
for c in a2.columns[1:]: a2[c] = a2[c].map(lambda v: f"{v:.4f}")
tabla(a2, size=9, titulo="Tabla 15. Autocorrelaciones del cambio logarítmico (rezagos 1 a 3)",
      fuente_nota="Interpretación: el nivel muestra correlaciones muy altas con sus tres rezagos (0.81–0.91); el cambio logarítmico, cercanas a cero (|r| ≤ 0.21).")
pregs(10)

# ======================= 11 y 12 =======================
doc.add_heading("11. Correlograma ACF en nivel", level=1)
figura("fig_11_acf_nivel", 6.1, "Función de autocorrelación (ACF) del precio de cierre de RKLB en nivel, 20 rezagos mensuales y banda de confianza al 95 %.",
       f"Las barras de los rezagos 1 a 4 superan la banda y el patrón decae lentamente: {n3(RS.A('acf_nivel',1,'ACF'))}, {n3(RS.A('acf_nivel',6,'ACF'))} (rezago 6), {n3(RS.A('acf_nivel',12,'ACF'))} (rezago 12), ≈ 0 en el rezago 15. Es la firma de una serie con tendencia y alta persistencia.")
ac = pd.DataFrame(R["acf_nivel"]).iloc[1:11][["Rezago", "ACF", "Banda sup.", "Fuera de bandas"]]
ac["Banda sup."] = ac["Banda sup."].map(lambda v: f"±{v:.3f}"); ac["ACF"] = ac["ACF"].map(lambda v: f"{v:.3f}"); ac["Fuera de bandas"] = ac["Fuera de bandas"].map({True: "Sí", False: "No"}); ac["Rezago"] = ac["Rezago"].astype(str)
tabla(ac, anchos=[1.0, 1.1, 1.3, 1.5], size=9, titulo="Tabla 16. ACF del nivel, rezagos 1 a 10")
pregs(11)
doc.add_heading("12. Correlograma PACF en nivel", level=1)
figura("fig_12_pacf_nivel", 6.1, "Función de autocorrelación parcial (PACF) del precio de cierre de RKLB en nivel, 20 rezagos mensuales.",
       f"Solo el rezago 1 es grande ({n3(RS.A('pacf_nivel',1,'PACF'))}); los rezagos 3 ({n3(RS.A('pacf_nivel',3,'PACF'))}) y 6 ({n3(RS.A('pacf_nivel',6,'PACF'))}) superan apenas la banda de ±{n3(R['banda_pacf_nivel'])}. La ACF alta en los rezagos 2–4 es, por tanto, un efecto en cadena del mes inmediato anterior.")
cm = pd.DataFrame({"Rezago": range(1, 9), "ACF": [RS.A('acf_nivel', k, 'ACF') for k in range(1, 9)], "PACF": [RS.A('pacf_nivel', k, 'PACF') for k in range(1, 9)]})
cm["|ACF| − |PACF|"] = cm["ACF"].abs() - cm["PACF"].abs()
for c in ["ACF", "PACF", "|ACF| − |PACF|"]: cm[c] = cm[c].map(lambda v: f"{v:.3f}")
cm["Rezago"] = cm["Rezago"].astype(str)
tabla(cm, anchos=[1.0, 1.2, 1.2, 1.6], size=9, titulo="Tabla 17. ACF vs. PACF del nivel (rezagos 1 a 8)",
      fuente_nota="Rezagos 2, 3 y 4: ACF alta (0.74–0.78) con PACF baja (−0.02 a 0.29).")
pregs(12)

# ======================= 13 =======================
doc.add_heading("13. ACF y PACF de la serie transformada (cambio logarítmico)", level=1)
para("Se eligió el **cambio logarítmico mensual** porque elimina la tendencia, expresa los movimientos en términos relativos y es aditivo en el tiempo (secciones 5 a 8).")
figura("fig_13_acf_pacf_comparacion", 6.5, "ACF y PACF de RKLB: nivel (fila superior) vs. cambio logarítmico mensual (fila inferior).",
       f"Fila superior: decaimiento lento y barras fuera de bandas. Fila inferior: todas las barras dentro de las bandas, sin patrón de decaimiento. Las más cercanas al límite son los rezagos 11 y 12 (ACF {n3(RS.A('acf_ret',11,'ACF'))} y {n3(RS.A('acf_ret',12,'ACF'))}), que no se pueden distinguir de ruido.")
figura("fig_13_acf_retorno", 5.6, "ACF del cambio logarítmico mensual de RKLB (%), 20 rezagos y banda de confianza al 95 %.",
       f"El rezago 1 es {n3(RS.A('acf_ret',1,'ACF'))} y el de mayor magnitud entre los tres primeros es el rezago 3 ({n3(RS.A('acf_ret',3,'ACF'))}), dentro de la banda de ±{n3(RS.A('acf_ret',3,'Banda sup.'))}.")
figura("fig_13_pacf_retorno", 5.6, "PACF del cambio logarítmico mensual de RKLB (%), 20 rezagos.",
       "La PACF repite la lectura de la ACF: no hay rezagos significativos; la dependencia lineal en los cambios logarítmicos mensuales es indistinguible de cero.")
cc = pd.DataFrame(R["comp_ac"]).iloc[:12]; cc["Rezago"] = cc["Rezago"].astype(int).astype(str)
for c in cc.columns[1:]: cc[c] = cc[c].map(lambda v: f"{v:.3f}")
tabla(cc, size=9, titulo="Tabla 18. ACF y PACF: nivel vs. cambio logarítmico (rezagos 1 a 12)",
      fuente_nota=f"Bandas de confianza al 95 %: nivel ≈ ±{n3(R['banda_nivel'])} (rezago 1); cambio log ≈ ±{n3(R['banda_ret'])}. Ljung-Box (10 rezagos): nivel p < 0.0001; cambio log p = {n3(R['lb_r_p'])}.")
pregs(13)

# ======================= 14 =======================
doc.add_heading("14. Observaciones atípicas y episodios", level=1)
figura("fig_14_episodios", 6.3, "Arriba: cambio logarítmico mensual de RKLB con meses atípicos (|z| > 1.5). Abajo: volatilidad móvil de 6 meses.",
       f"Un único mes supera 2 desviaciones (nov-2024). La volatilidad móvil pasa de {n1(R['vm_min'])} % (jul-2024) a {n1(R['vm_max'])} % (abr-2025) y vuelve a subir desde nov-2025: la varianza no es constante.")
ept = pd.DataFrame(R["ep"]); ept = ept[["Mes", "Cierre (USD)", "Cambio simple (%)", "Cambio log (%)", "z", "Volumen (M acc.)"]]
ept["Mes"] = ept["Mes"].map(fm) if ept["Mes"].str.len().max() > 8 else ept["Mes"]
ept["Mes"] = ept["Mes"].map(lambda s: fm(s + "-01") if re.match(r"^\d{4}-\d{2}$", s) else s)
for c in ept.columns[1:]: ept[c] = ept[c].map(lambda v: f"{v:+,.2f}" if c in ("Cambio simple (%)", "Cambio log (%)", "z") else f"{v:,.2f}")
tabla(ept, size=8.5, titulo="Tabla 19. Meses con |z| > 1.5 en el cambio logarítmico")
vo = pd.DataFrame(R["vol_anual"]); vo["anio"] = vo["anio"].astype(int).astype(str)
vo.columns = ["Año", "n", "Media (%)", "Desv. est. (%)", "Máx. (%)", "Mín. (%)"]; vo["n"] = vo["n"].astype(int).astype(str)
for c in vo.columns[2:]: vo[c] = vo[c].map(lambda v: f"{v:,.2f}")
tabla(vo, size=9, titulo="Tabla 20. Cambio logarítmico mensual por año (%)", fuente_nota="La desviación mensual pasa de ≈ 20 % (2023) a ≈ 27–30 % (2024–2026).")
pregs(14)

# ======================= 15 =======================
doc.add_heading("15. Conclusión diagnóstica integradora", level=1)
for q, a in RS.Q[15]:
    p = para(f"**{q}**", after=1); p.paragraph_format.keep_with_next = True
    para(a, indent=0.2)

# ======================= Fuentes =======================
doc.add_heading("Fuentes", level=1)
for s in [
    "**Datos:** archivo «Datos históricos RKLB» (CSV), precio mensual de Rocket Lab USA, Inc. (Nasdaq: RKLB), oct-2022 a oct-2026, con formato de exportación tipo Investing.com; proporcionado por el equipo; fecha de consulta: 2 de octubre de 2026. Enlace: [completar].",
    "**Indicaciones y rúbrica:** «Trabajo Final · Unidad 1 — Diagnóstico y preparación de una serie de tiempo» y Rúbrica de evaluación; Actividad 2 y Actividad 2-Parte II (Google Colab), Series de Tiempo Aplicadas a las Finanzas.",
    "**Software:** Python; pandas; NumPy; statsmodels (Seabold, S. y Perktold, J. (2010). *statsmodels: Econometric and statistical modeling with Python*, Proceedings of the 9th Python in Science Conference); matplotlib; seaborn; python-docx.",
    "No se consultó información externa para explicar los episodios atípicos; la base no permite identificar sus causas."]:
    p = doc.add_paragraph(style="List Bullet"); runs(p, s)
doc.add_heading("Entregables", level=2)
for s in ["`RKLB_Trabajo_Final_Unidad1.ipynb`: notebook ejecutado, con tablas, gráficas y respuestas.", "`Reporte_Ejecutivo_RKLB_Unidad1.docx`: este reporte.",
          "`Datos_historicos_RKLB.csv` (original) y `RKLB_base_limpia.csv` (limpia y ordenada): base de datos utilizada."]:
    p = doc.add_paragraph(style="List Bullet"); runs(p, s)

doc.save("Reporte_Ejecutivo_RKLB_Unidad1.docx")
print("Reporte guardado")
